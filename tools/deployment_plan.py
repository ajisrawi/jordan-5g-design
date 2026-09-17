#!/usr/bin/env python3
"""Deployment planner: splits the nominal design into 10 rollouts over 36 months.

Reads output/*.csv produced by generate_design.py. Writes deployment/*.md, deployment/*.csv,
deployment/Jordan_5G_Deployment_Plan.xlsx and data/rollout.js (map layer).
Atomic build unit = one hub with all its access rings (~40 sites), so fiber, DU hotel and
radio sites of a unit always land in the same rollout. Rural sites are grouped per governorate,
highway sites per route; a unit can never start before the backbone route that feeds it.
"""
import csv, json, math, os, datetime as dt
from collections import defaultdict, Counter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT, DEP = os.path.join(ROOT, "output"), os.path.join(ROOT, "deployment")
os.makedirs(DEP, exist_ok=True)
rd = lambda f: list(csv.DictReader(open(os.path.join(OUT, f), encoding="utf-8")))
sites, rings, hubs, venues = rd("sites.csv"), rd("fiber_rings.csv"), rd("hubs.csv"), rd("indoor_venues.csv")
summary = json.load(open(os.path.join(OUT, "summary.json")))["summary"]

START = dt.date(2027, 1, 1)                     # programme month 1 (assumption - shift as required)
NR = 10
QUOTA = [300, 450, 600, 700, 750, 750, 750, 700, 650, 10 ** 6]
def month_date(m): return dt.date(START.year + (START.month - 1 + m - 1) // 12, (START.month - 1 + m - 1) % 12 + 1, 1)
def win(r): return 4 + 3 * (r - 1), 6 + 3 * (r - 1)          # on-air window (programme months) of rollout r
def mlabel(m): return month_date(m).strftime("%b %Y")

# Backbone build schedule (route name in fiber_rings.csv -> rollout in which it must be lit)
BACKBONE_R = {"Amman-Zarqa-Mafraq-Jaber": 2, "Mafraq-Ramtha-Irbid": 2, "Irbid-Jerash-Amman (R35)": 2,
              "Desert Highway (R15) Amman-Aqaba": 3, "Dead Sea - Wadi Araba (R65) Amman-Aqaba": 4,
              "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an": 5, "Amman-Salt-Deir Alla": 5,
              "Jordan Valley (R65 north)": 6, "Irbid-Ajloun-Jerash": 6, "Karak-Qatraneh": 7, "Tafilah-Jurf": 7, "Amman-Azraq (R40)": 7,
              "Zarqa-Azraq-Safawi": 8, "Mafraq-Safawi-Ruwaished (R10)": 8, "Quweira-Wadi Rum": 8}
CLUSTER_MIN = {"Amman-Zarqa Metro": 1, "Irbid": 2, "Mafraq": 2, "Jerash": 2, "Aqaba": 3, "Ma'an": 3, "Dead Sea": 4, "South Shuna": 4, "Ghor Safi": 4,
               "Madaba": 5, "Karak": 5, "Tafilah": 5, "Wadi Musa / Petra": 5, "Mutah-Mazar": 5, "Salt": 5, "Deir Alla": 5,
               "North Shuna": 6, "Ajloun": 6, "Azraq": 7}
BONUS = {"Amman-Zarqa Metro": 0.5, "Irbid": 0.3, "Aqaba": 0.45, "Dead Sea": 0.6, "Wadi Musa / Petra": 0.6}
CENTRE = {"Amman-Zarqa Metro": (31.957, 35.912)}
W = {"DU": 3.0, "U": 2.0, "SU": 1.0}
POPD = {"DU": 25000, "U": 10000, "SU": 2500, "R": 150, "HW": 5}
ISD = {"DU": .375, "U": .45, "SU": .65, "R": 4.5, "HW": 7.0}
def km(a, b): return math.hypot((a[0] - b[0]) * 110.6, (a[1] - b[1]) * 111.3 * math.cos(math.radians(a[0])))

# ---------------------------------------------------------------- build units
hub_by = {h["hub_id"]: h for h in hubs}
units = defaultdict(lambda: dict(sites=[]))
for s in sites:
    if s["class"] in ("DU", "U", "SU"): k = s["hub"]
    elif s["class"] == "HW": k = "HW|" + s["zone"]
    else: k = "RUR|" + s["governorate"]
    units[k]["sites"].append(s)
final = {}
for k, u in units.items():
    if k.startswith("RUR|"):                                   # chunk rural governorates into <= 40-site work packages, north to south
        ss = sorted(u["sites"], key=lambda s: -float(s["lat"]))
        for i in range(0, len(ss), 40): final[f"{k}|{i // 40 + 1}"] = dict(sites=ss[i:i + 40])
    else: final[k] = u
for k, u in final.items():
    ss = u["sites"]; u["n"] = len(ss)
    if k.startswith("HUB"):
        cl = hub_by[k]["cluster"]; pos = (float(hub_by[k]["lat"]), float(hub_by[k]["lon"]))
        lat0 = sum(float(s["lat"]) for s in sites if s["cluster"] == cl) / sum(1 for s in sites if s["cluster"] == cl)
        lon0 = sum(float(s["lon"]) for s in sites if s["cluster"] == cl) / sum(1 for s in sites if s["cluster"] == cl)
        d = km(pos, CENTRE.get(cl, (lat0, lon0)))
        u.update(kind="Hub cluster", cluster=cl, min_r=CLUSTER_MIN.get(cl, 5),
                 score=sum(W[s["class"]] for s in ss) / len(ss) + BONUS.get(cl, 0) - 0.02 * d)
    elif k.startswith("HW|"):
        route = k[3:]; nat = BACKBONE_R.get(route, 8) <= 4
        u.update(kind="Highway corridor", cluster=route, min_r=min(NR, BACKBONE_R.get(route, 8) + 1), score=5.0 if nat else 4.5)   # ride on the backbone as soon as it is lit
    else:
        u.update(kind="Rural package", cluster=k.split("|")[1] + " governorate (rural)", min_r=7 + (int(k.split("|")[2]) - 1) % 4, score=4.0)  # rural spread over RO-07..10

assign = {}
for r in range(1, NR + 1):
    tot = 0
    for k, u in sorted(final.items(), key=lambda kv: -kv[1]["score"]):
        if k in assign or u["min_r"] > r: continue
        if tot + u["n"] > QUOTA[r - 1] + 20: continue
        assign[k] = r; tot += u["n"]
for k in final:
    assign.setdefault(k, NR)
site_r = {s["site_id"]: assign[k] for k, u in final.items() for s in u["sites"]}
hub_r = {k: r for k, r in assign.items() if k.startswith("HUB")}

# ---------------------------------------------------------------- per-rollout quantities
ring_km = defaultdict(float)
for g in rings:
    if g["parent"] in hub_r: ring_km[hub_r[g["parent"]]] += float(g["route_km"])
agg_by_pop = defaultdict(float)
for g in rings:
    if g["tier"].startswith("aggregation"): agg_by_pop[g["parent"]] += float(g["route_km"])
hubs_per_pop = Counter(h["pop"] for h in hubs)
agg_km = defaultdict(float)
for h in hubs: agg_km[hub_r[h["hub_id"]]] += agg_by_pop[h["pop"]] / hubs_per_pop[h["pop"]]
bb_km = defaultdict(float); bb_routes = defaultdict(list)
for g in rings:
    if g["tier"].startswith("DWDM"):
        r = BACKBONE_R.get(g["ring_id"], 8); bb_km[r] += float(g["route_km"]); bb_routes[r].append(f"{g['ring_id']} ({g['route_km']} km)")
bb_km[1] += summary["metro_core_km"]; bb_routes[1].insert(0, f"Amman metro core ring ({summary['metro_core_km']} km)")
n_spur = sum(1 for s in sites if "spur" in s["backhaul"]); spur_each = summary["spur_km"] / max(n_spur, 1)

site_pos = {s["site_id"]: (float(s["lat"]), float(s["lon"])) for s in sites}
ven_r = {}
for v in venues:
    p = (float(v["lat"]), float(v["lon"]))
    near = min(sites, key=lambda s: km(p, site_pos[s["site_id"]]))
    ven_r[v["venue"]] = min(NR, max(2, site_r[near["site_id"]] + 1))

RO = []
cum_sites = cum_pop = cum_area = 0
tot_pop = sum(POPD[s["class"]] * 0.866 * ISD[s["class"]] ** 2 for s in sites)
tot_1g = sum(0.866 * ISD[s["class"]] ** 2 for s in sites if s["class"] in ("DU", "U", "SU"))
for r in range(1, NR + 1):
    ss = [s for s in sites if site_r[s["site_id"]] == r]
    c = Counter(s["class"] for s in ss); t = Counter(s["site_type"] for s in ss)
    sec = sum(int(s["sectors"]) for s in ss); sec_city = sum(int(s["sectors"]) for s in ss if s["class"] in ("DU", "U", "SU"))
    sec_rur = sum(int(s["sectors"]) for s in ss if s["class"] == "R")
    my_hubs = [h for h in hubs if hub_r[h["hub_id"]] == r]
    cran_sites = sum(1 for s in ss if s["hub"] in hub_by and hub_by[s["hub"]]["type"].startswith("C-RAN"))
    pop = sum(POPD[s["class"]] * 0.866 * ISD[s["class"]] ** 2 for s in ss)
    area = sum(0.866 * ISD[s["class"]] ** 2 for s in ss if s["class"] in ("DU", "U", "SU"))
    cum_sites += len(ss); cum_pop += pop * summary["pop_in_zones"] / tot_pop; cum_area += area * summary["geo_1g_km2"] / tot_1g
    vv = [v for v in venues if ven_r[v["venue"]] == r]
    a, b = win(r)
    RO.append(dict(
        r=r, id=f"RO-{r:02d}", m_from=a, m_to=b, window=f"{mlabel(a)} – {mlabel(b)}", sites=len(ss), cls=dict(c), types=dict(t), sectors=sec,
        aau64=sec_city, aau8=sec_rur, passive=sec, rru_low=sec, rru_mid=sec, n258=(480 if r in (9, 10) else 0),
        du_cran=math.ceil(cran_sites * 0.75), du_dran=len(ss) - cran_sites, csr=len(ss) - cran_sites, fhgw=cran_sites, hubs=len(my_hubs), hub_routers=2 * len(my_hubs),
        mw=sum(1 for s in ss if s["backhaul"].startswith("Micro")), access_km=round(ring_km[r]), agg_km=round(agg_km[r]), bb_km=round(bb_km[r]),
        spur_km=round(sum(spur_each for s in ss if "spur" in s["backhaul"])), bb_routes=bb_routes[r],
        rooftop=t.get("Rooftop", 0), mono=t.get("Monopole / greenfield", 0), tower=t.get("Lattice tower", 0),
        clusters=Counter(final[k]["cluster"] for k, x in assign.items() if x == r for _ in final[k]["sites"]).most_common(),
        units=sorted([(k, final[k]["kind"], final[k]["cluster"], final[k]["n"]) for k, x in assign.items() if x == r], key=lambda q: -q[3]),
        gov=Counter(s["governorate"] for s in ss).most_common(), venues=[(v["venue"], v["type"], int(v["n78_pRRU_4T4R"]), int(v["n258_mmWave_heads"])) for v in vv],
        prru=sum(int(v["n78_pRRU_4T4R"]) for v in vv), vmmw=sum(int(v["n258_mmWave_heads"]) for v in vv),
        cum_sites=cum_sites, cum_pop_pct=round(100 * cum_pop / summary["pop_total"], 1), cum_1g_km2=round(cum_area),
        rate=round(len(ss) / 3)))

# ---------------------------------------------------------------- CSV / JS / XLSX
with open(os.path.join(DEP, "rollout_sites.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["rollout", "on_air_window", "site_id", "class", "site_type", "cluster", "zone", "governorate", "lat", "lon", "antenna_height_agl_m", "sectors", "hub", "ring", "backhaul", "is_hub"])
    for s in sorted(sites, key=lambda s: (site_r[s["site_id"]], s["hub"], s["ring"], s["site_id"])):
        r = site_r[s["site_id"]]; w.writerow([f"RO-{r:02d}", RO[r - 1]["window"], s["site_id"], s["class"], s["site_type"], s["cluster"], s["zone"], s["governorate"], s["lat"], s["lon"],
                                              s["antenna_height_agl_m"], s["sectors"], s["hub"], s["ring"], s["backhaul"], s["is_hub"]])
BOQ_ROWS = [("Macro sites", "sites"), ("– rooftop", "rooftop"), ("– monopole / greenfield", "mono"), ("– lattice tower 45 m", "tower"), ("Sectors", "sectors"),
            ("n78 64T64R AAU (3×100 MHz)", "aau64"), ("n78 8T8R radio + antenna (rural)", "aau8"), ("Passive multiband antenna 2L4H + RET", "passive"),
            ("n28 RRU", "rru_low"), ("n1/n3 dual-band RRU", "rru_mid"), ("n258 street small cells", "n258"), ("DU (pooled, C-RAN hub)", "du_cran"), ("DU (at site, D-RAN)", "du_dran"),
            ("Cell-site router 25GE (D-RAN sites)", "csr"), ("Fronthaul gateway / passive WDM (C-RAN sites)", "fhgw"), ("Hubs (DU hotel / pre-agg) commissioned", "hubs"), ("Hub aggregation routers 100GE", "hub_routers"), ("Microwave hops", "mw"),
            ("Access-ring fiber (route km)", "access_km"), ("Aggregation-ring fiber (route km)", "agg_km"), ("Backbone / metro-core fiber lit (route km)", "bb_km"),
            ("Backbone add/drop spurs (km)", "spur_km"), ("Indoor venues", None), ("Indoor n78 pRRU", "prru"), ("Indoor n258 heads", "vmmw")]
def val(x, k): return len(x["venues"]) if k is None else x[k]
with open(os.path.join(DEP, "rollout_boq.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["item"] + [x["id"] for x in RO] + ["total"])
    for name, k in BOQ_ROWS: w.writerow([name] + [val(x, k) for x in RO] + [sum(val(x, k) for x in RO)])
PIPE = [("Nominal release, TSSR surveys, site acquisition & permits", -6, -3), ("Detailed RF / transport design freeze, material call-off", -5, -3),
        ("Civil works, power, fiber build to hubs and rings", -4, -1), ("Hub / DU-hotel and router commissioning", -2, 0),
        ("Radio install, commissioning, integration to 5GC", -1, 2), ("Single-site verification (SSV)", 0, 2),
        ("Cluster optimisation & acceptance (1 Gbps drive test)", 1, 4), ("Commercial launch of rollout area / hand-over to operations", 3, 4)]
sched = []
for x in RO:
    for name, a, b in PIPE:
        m0, m1 = max(1, x["m_from"] + a), min(36, x["m_from"] + b)
        sched.append([x["id"], name, m0, m1, mlabel(m0), mlabel(m1)])
with open(os.path.join(DEP, "rollout_schedule.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["rollout", "activity", "start_month", "end_month", "start", "end"]); w.writerows(sched)
with open(os.path.join(DEP, "rollout_venues.csv"), "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f); w.writerow(["rollout", "venue", "type", "n78_pRRU", "n258_heads"])
    for x in RO:
        for v in x["venues"]: w.writerow([x["id"]] + list(v))
COL = ["#e6194b", "#f58231", "#ffe119", "#bfef45", "#3cb44b", "#42d4f4", "#4363d8", "#911eb4", "#f032e6", "#a9a9a9"]
with open(os.path.join(ROOT, "data", "rollout.js"), "w", encoding="utf-8") as f:
    f.write("window.ROLLOUT=" + json.dumps(dict(site=site_r, hub=hub_r, venue=ven_r, colors=COL, backbone=BACKBONE_R,
            meta=[dict(id=x["id"], window=x["window"], sites=x["sites"], cum_sites=x["cum_sites"], cum_pop=x["cum_pop_pct"], cum_1g=x["cum_1g_km2"],
                       top=", ".join(f"{c} ({n})" for c, n in x["clusters"][:4])) for x in RO]), separators=(",", ":"), ensure_ascii=False) + ";")

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb = Workbook(); hd = PatternFill("solid", fgColor="1F3A5F"); hf = Font(bold=True, color="FFFFFF")
def sheet(name, header, rows, widths=None, first=False):
    ws = wb.active if first else wb.create_sheet(); ws.title = name; ws.append(header)
    for c in ws[1]: c.fill, c.font, c.alignment = hd, hf, Alignment(wrap_text=True, vertical="center")
    for r in rows: ws.append(r)
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    for i, _ in enumerate(header, 1): ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 16)
    return ws
sheet("Rollout summary", ["Rollout", "On-air window", "Sites", "DU", "U", "SU", "Rural", "Highway", "Sectors", "Hubs", "Access fiber km", "Agg fiber km", "Backbone km", "Indoor venues", "Cum. sites", "Cum. pop. covered %", "Cum. 1 Gbps km²", "Sites / month", "Main areas"],
      [[x["id"], x["window"], x["sites"]] + [x["cls"].get(c, 0) for c in ("DU", "U", "SU", "R", "HW")] + [x["sectors"], x["hubs"], x["access_km"], x["agg_km"], x["bb_km"], len(x["venues"]), x["cum_sites"], x["cum_pop_pct"], x["cum_1g_km2"], x["rate"],
        ", ".join(f"{c} ({n})" for c, n in x["clusters"][:5])] for x in RO], {2: 22, 19: 80}, first=True)
sheet("BoQ by rollout", ["Item"] + [x["id"] for x in RO] + ["Total"], [[n] + [val(x, k) for x in RO] + [sum(val(x, k) for x in RO)] for n, k in BOQ_ROWS], {1: 44})
ws = sheet("Schedule (Gantt)", ["Rollout", "Activity"] + [f"M{m}\n{mlabel(m)}" for m in range(1, 37)], [], {1: 9, 2: 58})
fills = [PatternFill("solid", fgColor=c[1:].upper()) for c in COL]
for row in sched:
    ws.append([row[0], row[1]] + [""] * 36); rr = ws.max_row
    for m in range(row[2], row[3] + 1): ws.cell(rr, 2 + m).fill = fills[int(row[0][3:]) - 1]
for i in range(3, 39): ws.column_dimensions[get_column_letter(i)].width = 5.5
ws.row_dimensions[1].height = 32
sheet("Work packages", ["Rollout", "Work package", "Type", "Cluster / route", "Sites"], [[x["id"], u[0], u[1], u[2], u[3]] for x in RO for u in x["units"]], {2: 26, 3: 18, 4: 50})
sheet("Sites", ["Rollout", "Site", "Class", "Type", "Cluster", "Governorate", "Lat", "Lon", "Ant. height m", "Sectors", "Hub", "Ring", "Backhaul"],
      [[f"RO-{site_r[s['site_id']]:02d}", s["site_id"], s["class"], s["site_type"], s["cluster"], s["governorate"], float(s["lat"]), float(s["lon"]), int(s["antenna_height_agl_m"]), int(s["sectors"]), s["hub"], s["ring"], s["backhaul"]]
       for s in sorted(sites, key=lambda s: (site_r[s["site_id"]], s["site_id"]))], {2: 16, 5: 22, 11: 14, 12: 18, 13: 34})
sheet("Indoor", ["Rollout", "Venue", "Type", "n78 pRRU", "n258 heads"], [[x["id"]] + list(v) for x in RO for v in x["venues"]], {2: 44})
wb.save(os.path.join(DEP, "Jordan_5G_Deployment_Plan.xlsx"))

json.dump(RO, open(os.path.join(DEP, "rollouts.json"), "w", encoding="utf-8"), indent=1, ensure_ascii=False)
for x in RO:
    print(x["id"], x["window"], x["sites"], x["cls"], "hubs", x["hubs"], "acc", x["access_km"], "agg", x["agg_km"], "bb", x["bb_km"], "ven", len(x["venues"]), "cum", x["cum_sites"], x["cum_pop_pct"], x["cum_1g_km2"], [c for c, _ in x["clusters"][:6]])
