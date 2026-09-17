#!/usr/bin/env python3
"""Low-level (equipment) design, vendor-neutral: RAN equipment + power, 5G core sizing, transport ports / DWDM / microwave.

Reads output/*.csv + summary.json (+ DEM for microwave line-of-sight). Writes lld/*.md, lld/*.csv, lld/Jordan_5G_LLD.xlsx.
All unit capacities below are typical 2026 carrier-grade figures, deliberately vendor-neutral: replace them with the
selected vendor's datasheet values and re-run.
"""
import csv, json, math, os, sys
from collections import Counter, defaultdict
import numpy as np
sys.path.insert(0, os.path.dirname(__file__))
from terrain import Terrain

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT, LLD = os.path.join(ROOT, "output"), os.path.join(ROOT, "lld")
os.makedirs(LLD, exist_ok=True)
rd = lambda f: list(csv.DictReader(open(os.path.join(OUT, f), encoding="utf-8")))
sites, rings, hubs = rd("sites.csv"), rd("fiber_rings.csv"), rd("hubs.csv")
J = json.load(open(os.path.join(OUT, "summary.json"))); S = J["summary"]
OFFER = {r["cls"]: r["offered_mbps"] for r in J["capacity"]}; OFFER["HW"] = 30        # busy-hour Mbps per sector
fm = lambda n: f"{n:,.0f}"
def wcsv(name, header, rows):
    with open(os.path.join(LLD, name), "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f); w.writerow(header); w.writerows(rows)
def wmd(name, txt):
    open(os.path.join(LLD, name), "w", encoding="utf-8").write(txt.strip() + "\n"); print("wrote", name)
def table(header, rows):
    return "| " + " | ".join(header) + " |\n|" + "---|" * len(header) + "\n" + "".join("| " + " | ".join(str(c) for c in r) + " |\n" for r in rows)

# =====================================================================================================
# 1. RAN
# =====================================================================================================
# unit: (average W, peak W)
PWR = {"AAU64": (900, 1300), "R8T": (750, 1050), "RRU28_2T": (350, 500), "RRU28_4T": (550, 750), "RRU_MID": (700, 950), "RRU_N3": (450, 620),
       "DU": (350, 500), "CSR": (120, 150), "FHGW": (60, 80), "AUX": (150, 200)}
hub_by = {h["hub_id"]: h for h in hubs}
def is_cran(s): return s["hub"] in hub_by and hub_by[s["hub"]]["type"].startswith("C-RAN") and s["class"] in ("DU", "U", "SU")
def site_equip(s):
    c, n = s["class"], int(s["sectors"]); cran = is_cran(s)
    e = Counter()
    if c in ("DU", "U", "SU"):
        e["AAU64"] = n; e["RRU28_2T"] = n; e["RRU_MID" if c != "SU" else "RRU_N3"] = n
        cells = n * (3 + (3 if c != "SU" else 2))
    elif c == "R": e["R8T"] = n; e["RRU28_4T"] = n; e["RRU_MID"] = n; cells = n * 6
    else: e["RRU28_4T"] = n; e["RRU_N3"] = n; cells = n * 2
    e["ANT"] = n; e["AUX"] = 1
    if cran: e["FHGW"] = 1
    else: e["DU"] = 1; e["CSR"] = 1
    avg = sum(PWR[k][0] * v for k, v in e.items() if k in PWR); peak = sum(PWR[k][1] * v for k, v in e.items() if k in PWR)
    hours = 4 if c in ("DU", "U", "SU") else 8
    batt_kwh = avg / 1000 * hours / (0.8 * 0.95)                       # 80 % DoD, 95 % conversion
    mods = math.ceil(batt_kwh / 4.8)                                   # 48 V 100 Ah Li-ion module = 4.8 kWh
    rect_kw = math.ceil((peak * 1.1 + mods * 4.8 * 1000 * 0.2) / 3000) + 1   # 3 kW modules, 0.2C recharge, N+1
    offgrid = c in ("R", "HW") and s["backhaul"].startswith("Micro")
    pv = round(avg / 1000 * 24 * 0.6 / (5.5 * 0.75), 1) if offgrid else 0   # traffic-adaptive sleep 0.6, 5.5 peak-sun-hours, 75 % system efficiency
    return dict(e=e, cells=cells, avg=avg, peak=peak, batt=mods * 4.8, mods=mods, rect=rect_kw * 3, kva=math.ceil(peak * 1.25 / 900 / 5) * 5, pv=pv, cran=cran, offgrid=offgrid)
SE = {s["site_id"]: site_equip(s) for s in sites}
wcsv("ran_site_equipment.csv", ["site_id", "class", "site_type", "architecture", "sectors", "n78_64T64R_AAU", "n78_8T8R_radio", "n28_RRU", "n1n3_or_n3_RRU", "passive_antenna_2L4H",
     "DU_at_site", "CSR", "fronthaul_gateway", "NR_cells", "avg_load_W", "peak_load_W", "rectifier_kW", "battery_kWh", "battery_modules_4.8kWh", "grid_kVA", "solar_PV_kWp"],
     [[s["site_id"], s["class"], s["site_type"], "C-RAN" if x["cran"] else "D-RAN", s["sectors"], x["e"]["AAU64"], x["e"]["R8T"], x["e"]["RRU28_2T"] + x["e"]["RRU28_4T"],
       x["e"]["RRU_MID"] + x["e"]["RRU_N3"], x["e"]["ANT"], x["e"]["DU"], x["e"]["CSR"], x["e"]["FHGW"], x["cells"], x["avg"], x["peak"], x["rect"], round(x["batt"], 1), x["mods"],
       0 if x["offgrid"] else x["kva"], x["pv"]] for s in sites for x in [SE[s["site_id"]]]])
# typical configurations
TYP = {}
for s in sites:
    k = (s["class"], "C-RAN" if SE[s["site_id"]]["cran"] else "D-RAN", s["sectors"])
    TYP.setdefault(k, [0, SE[s["site_id"]]]); TYP[k][0] += 1
typ_rows = [[f"{k[0]} {k[1]} {k[2]}-sector", n, x["cells"], fm(x["avg"]), fm(x["peak"]), x["rect"], round(x["batt"], 1), x["kva"]] for k, (n, x) in sorted(TYP.items(), key=lambda q: -q[1][0]) if n >= 20]

# hub baseband pools
POOL_GAIN = 0.75                  # pooled DUs per C-RAN site (25 % statistical pooling gain vs 1 DU per site)
DU_PER_RACK, DU_W = 12, 500
hub_rows, tot_du_pool = [], 0
ring_cnt = Counter(g["parent"] for g in rings if g["tier"].startswith("access"))
hub_sites = defaultdict(list)
for s in sites:
    if s["hub"]: hub_sites[s["hub"]].append(s)
HUBX = {}
for h in hubs:
    ss = hub_sites[h["hub_id"]]; cran = h["type"].startswith("C-RAN")
    ndu = math.ceil(len(ss) * POOL_GAIN) if cran else 0; tot_du_pool += ndu
    bh = sum(OFFER[s["class"]] * int(s["sectors"]) for s in ss) / 1000          # Gbps busy hour
    peak = bh * 1.5 + 5.2                                                         # 1.5 burst factor + one sector at full 300 MHz peak
    it_kw = (ndu * DU_W + 2 * 1800 + 600) / 1000; fac_kw = it_kw * 1.4
    HUBX[h["hub_id"]] = dict(n=len(ss), cran=cran, ndu=ndu, rings=ring_cnt[h["hub_id"]], bh=bh, peak=peak)
    hub_rows.append([h["hub_id"], h["cluster"], h["type"].split(" ")[0], len(ss), ring_cnt[h["hub_id"]], ndu, math.ceil(ndu / DU_PER_RACK) + 2, round(it_kw, 1), round(fac_kw, 1),
                     math.ceil(fac_kw * 1.25 / 10) * 10, round(bh, 1), round(peak, 1)])
wcsv("ran_hub_baseband.csv", ["hub_id", "cluster", "type", "sites", "access_rings", "pooled_DUs", "racks_incl_transport_power", "IT_load_kW", "facility_load_kW", "grid_kVA", "busy_hour_Gbps", "peak_Gbps"], hub_rows)
EQ = Counter()
for x in SE.values(): EQ.update(x["e"])
cells_total = sum(x["cells"] for x in SE.values())
n_off = sum(1 for x in SE.values() if x["offgrid"])

# =====================================================================================================
# 2. CORE
# =====================================================================================================
SUBS = S["subs_total"]; BH_GBPS = S["bh_total_gbps"]
A = dict(reg=0.85, sess=1.6, volte=0.6, bhca=1.0, hold=90, peak=1.2, util=0.7, upf_srv=180, vcpu_srv=96, srv_rack=16)
ev = dict(registration=1.2, pdu_session=3.0, service_request=12.0, handover=8.0, paging=4.0)            # per subscriber per busy hour
tps = {k: SUBS * A["reg"] * v / 3600 for k, v in ev.items()}
NF = [  # name, driver, per-instance capacity, vCPU per instance, note
    ("AMF", SUBS * A["reg"], 2.0e6, 640, "registered subs"), ("SMF", SUBS * A["reg"] * A["sess"], 3.0e6, 720, "PDU sessions"),
    ("AUSF + UDM", SUBS, 5.0e6, 480, "provisioned subs"), ("UDR (subscriber DB)", SUBS, 10.0e6, 960, "provisioned subs, 3 replicas"),
    ("PCF + BSF", SUBS * A["reg"] * A["sess"], 3.0e6, 560, "policy sessions"), ("CHF (charging)", SUBS * A["reg"] * A["sess"], 4.0e6, 640, "charging sessions"),
    ("NRF + NSSF + SCP", sum(tps.values()) * 6, 60000, 320, "SBI transactions/s (≈ 6 per procedure)"), ("NEF + NWDAF", SUBS, 10.0e6, 480, "exposure / analytics"),
    ("IMS (P/I/S-CSCF, TAS, MRF, SBC)", SUBS * A["volte"], 1.5e6, 1100, "VoNR subs"), ("SEPP + roaming / interconnect GW", SUBS * 0.05, 1.0e6, 240, "roamers"),
    ("SMSF + CBCF (SMS, public warning)", SUBS, 10.0e6, 200, "subs"), ("OSS, observability, LI mediation, security", SUBS, 5.0e6, 900, "platform")]
nf_rows, vcpu_site = [], 0
for n, drv, cap, vc, note in NF:
    inst = max(2, math.ceil(drv / cap / A["util"]) + 1)                 # N+1, 70 % design utilisation, minimum 2
    nf_rows.append([n, fm(drv), note, fm(cap), inst, inst * vc]); vcpu_site += inst * vc
erl = SUBS * A["volte"] * A["bhca"] * A["hold"] / 3600
cp_srv = math.ceil(vcpu_site * 1.25 / A["vcpu_srv"])                    # 25 % platform overhead (CaaS, storage, monitoring)
wcsv("core_nf_sizing.csv", ["network_function", "dimensioning_driver", "driver_unit", "capacity_per_instance", "instances_per_DC_N+1", "vCPU_per_DC"], nf_rows)
SPLIT = [("DC1 Amman West", 0.275, True), ("DC2 Amman East", 0.275, True), ("Amman Centre edge", 0.08, False), ("Zarqa edge", 0.07, False), ("Irbid regional DC", 0.15, False),
         ("Aqaba regional DC", 0.08, False), ("Regional PoPs (Karak, Ma'an, Mafraq)", 0.07, False)]
core_rows = []
for n, sh, cp in SPLIT:
    gb = BH_GBPS * sh * A["peak"]; upf = math.ceil(gb / A["util"] / A["upf_srv"]) + 1
    srv = upf + (cp_srv if cp else 12 if "regional DC" in n else 6)
    racks = math.ceil(srv / A["srv_rack"]) + 2
    core_rows.append([n, f"{sh * 100:.1f} %", round(gb), upf, cp_srv if cp else 0, srv, racks, round(srv * 0.75 + 2 * 6), math.ceil(gb / A["util"] / 400) * 2])
wcsv("core_site_sizing.csv", ["site", "user_plane_share", "peak_Gbps", "UPF_servers_N+1", "control_plane_servers", "total_servers", "racks", "IT_load_kW", "400GE_ports_to_transport"], core_rows)

# =====================================================================================================
# 3. TRANSPORT
# =====================================================================================================
port_rows = []
for h in hubs:
    x = HUBX[h["hub_id"]]
    p25 = x["rings"] + math.ceil(x["ndu"] / 2) + 2                        # per router: one end of every ring, half the DUs, 2 spare
    up = 2 if x["peak"] < 60 else 4
    cls = "HUB-R1 (24×25GE + 4×100GE)" if p25 <= 24 and up <= 4 else "HUB-R2 (48×25GE + 8×100GE)"
    port_rows.append([h["hub_id"], h["pop"], x["n"], x["rings"], x["ndu"], p25, up, cls, round(x["bh"], 1), round(x["peak"], 1), f"{100 * x['peak'] / (up * 100):.0f} %"])
wcsv("transport_hub_ports.csv", ["hub_id", "parent_pop", "sites", "access_rings", "pooled_DUs", "25GE_ports_per_router", "100GE_uplinks_per_router", "router_class_x2", "busy_hour_Gbps", "peak_Gbps", "uplink_utilisation_at_peak"], port_rows)
hub_cls = Counter(r[7] for r in port_rows)
# access ring check
acc = [g for g in rings if g["tier"].startswith("access")]
ring_load = defaultdict(float)
for s in sites:
    if s["ring"]: ring_load[s["ring"]] += OFFER[s["class"]] * int(s["sectors"]) / 1000
dran_rings = [ring_load[g["ring_id"]] * 1.5 + 5.2 for g in acc if not g["tier"].startswith("access 25GE (eCPRI")]
# PoP routers
agg = [g for g in rings if g["tier"].startswith("aggregation")]
pop_rings = Counter(g["parent"] for g in agg); pop_bh = defaultdict(float)
for h in hubs: pop_bh[h["pop"]] += HUBX[h["hub_id"]]["bh"]
pop_rows = []
for p, n in sorted(pop_rings.items(), key=lambda q: -pop_bh[q[0]]):
    pk = pop_bh[p] * 1.3; up400 = max(2, math.ceil(pk / 0.5 / 400))
    pop_rows.append([p, n, n, up400, "POP-R2 (36×100GE + 8×400GE)" if n > 8 or up400 > 4 else "POP-R1 (16×100GE + 4×400GE)", round(pop_bh[p], 1), round(pk, 1)])
wcsv("transport_pop_ports.csv", ["pop", "aggregation_rings", "100GE_ports_per_router", "400GE_uplinks_per_router", "router_class_x2", "busy_hour_Gbps", "peak_Gbps"], pop_rows)
agg_rows = [[g["ring_id"], g["parent"], g["nodes"], g["route_km"], round(float(g["route_km"]) / (int(g["nodes"]) + 1), 1),
             "100G-LR4 (10 km)" if float(g["route_km"]) / (int(g["nodes"]) + 1) <= 8 else "100G-ER4 (40 km)" if float(g["route_km"]) / (int(g["nodes"]) + 1) <= 32 else "100G-ZR coherent (80 km)"] for g in agg]
wcsv("transport_agg_ring_optics.csv", ["ring_id", "pop", "hubs", "route_km", "mean_span_km", "optic"], agg_rows)

# DWDM wavelength plan
cl_bh = defaultdict(float)
for s in sites: cl_bh[s["cluster"]] += OFFER[s["class"]] * int(s["sectors"]) / 1000
gov_rural = defaultdict(float)
for s in sites:
    if s["class"] in ("R", "HW"): gov_rural[s["governorate"]] += OFFER[s["class"]] * int(s["sectors"]) / 1000
HOME = {"Irbid": "Irbid-Jerash-Amman (R35)", "Jerash": "Irbid-Jerash-Amman (R35)", "Mafraq": "Amman-Zarqa-Mafraq-Jaber", "Ajloun": "Irbid-Ajloun-Jerash", "North Shuna": "Jordan Valley (R65 north)",
        "Aqaba": "Desert Highway (R15) Amman-Aqaba", "Ma'an": "Desert Highway (R15) Amman-Aqaba", "Dead Sea": "Dead Sea - Wadi Araba (R65) Amman-Aqaba", "South Shuna": "Dead Sea - Wadi Araba (R65) Amman-Aqaba",
        "Ghor Safi": "Dead Sea - Wadi Araba (R65) Amman-Aqaba", "Madaba": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an", "Karak": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an",
        "Tafilah": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an", "Wadi Musa / Petra": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an", "Mutah-Mazar": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an",
        "Salt": "Amman-Salt-Deir Alla", "Deir Alla": "Amman-Salt-Deir Alla", "Azraq": "Amman-Azraq (R40)"}
GOV_HOME = {"Irbid": "Irbid-Jerash-Amman (R35)", "Ajloun": "Irbid-Ajloun-Jerash", "Jerash": "Irbid-Jerash-Amman (R35)", "Mafraq": "Amman-Zarqa-Mafraq-Jaber", "Balqa": "Amman-Salt-Deir Alla",
            "Karak": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an", "Tafilah": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an", "Madaba": "King's Highway Madaba-Karak-Tafilah-Petra-Ma'an",
            "Ma'an": "Desert Highway (R15) Amman-Aqaba", "Aqaba": "Desert Highway (R15) Amman-Aqaba"}
work = defaultdict(float)
for c, r in HOME.items(): work[r] += cl_bh.get(c, 0)
for g, r in GOV_HOME.items(): work[r] += gov_rural.get(g, 0)
INTL = BH_GBPS * 0.4 * 0.4                                                # 40 % of internet via Aqaba landing, 60 % served from in-country caches
RINGS = {"Northern ring": ["Amman-Zarqa-Mafraq-Jaber", "Mafraq-Ramtha-Irbid", "Irbid-Jerash-Amman (R35)"],
         "Southern ring": ["Desert Highway (R15) Amman-Aqaba", "Dead Sea - Wadi Araba (R65) Amman-Aqaba"]}
feeds = {"Northern ring": ["Irbid-Ajloun-Jerash", "Jordan Valley (R65 north)"], "Southern ring": ["King's Highway Madaba-Karak-Tafilah-Petra-Ma'an"]}
ring_dem = {}
for rn, members in RINGS.items():
    ring_dem[rn] = sum(work[m] for m in members) + sum(work[f] * 0.5 for f in feeds[rn]) + (INTL if rn == "Southern ring" else 0)
dw_rows = []
for g in rings:
    if not g["tier"].startswith("DWDM"): continue
    r = g["ring_id"]; kmr = float(g["route_km"]); rn = next((k for k, v in RINGS.items() if r in v), None)
    dem = (ring_dem[rn] if rn else work.get(r, 0)) * 1.3 * 1.5                # 30 % peak, 50 % 3-year growth headroom
    small = rn is None and dem < 160                                     # lightly loaded spurs start on 100G wavelengths
    lam = 2 if small else max(2, math.ceil(dem / 400))
    dw_rows.append([r, rn or "spur / regional", round(kmr), round(dem), lam, f"{lam} × {'100G' if small else '400G'} (of 64 × 75 GHz C-band slots)", max(0, math.ceil(kmr / 80) - 1),
                    "ROADM (2-degree, 3 at junctions)" if rn else "FOADM / terminal", ("100G coherent" if small else "400G ZR+/OTN coherent, FlexO" if kmr > 40 else "400G ZR")])
wcsv("dwdm_wavelength_plan.csv", ["route", "ring", "route_km", "design_demand_Gbps", "wavelengths_day1", "channel_plan", "inline_amplifier_sites_80km", "node_type", "line_optics"], dw_rows)

# Microwave hops with DEM line-of-sight
DEM = Terrain(os.path.join(ROOT, "data", "raw", "dem"))
fib = [s for s in sites if s["backhaul"].startswith("Fiber")]
FL = np.array([[float(s["lat"]), float(s["lon"])] for s in fib])
BANDS = [  # max km, name, f GHz, dish m, tx dBm, threshold dBm at design modulation, capacity, k, alpha (ITU-R P.838 horizontal pol.)
    (3.0, "E-band 80 GHz 1+0, 2 GHz channel", 80, 0.6, 15, -58, "10 Gbps", 1.10, 0.72),
    (7.0, "E-band + 18 GHz multiband (carrier-bonded)", 18, 0.9, 20, -62, "10 Gbps clear-sky / 1 Gbps guaranteed", 0.0708, 1.08),
    (16.0, "18 GHz XPIC 2+0, 2 × 112 MHz, 2048-QAM", 18, 1.2, 22, -60, "2.0 Gbps", 0.0708, 1.08),
    (32.0, "11 GHz XPIC 2+0, 2 × 80 MHz, 1024-QAM", 11, 1.8, 26, -62, "1.3 Gbps", 0.0177, 1.20),
    (999, "7 GHz XPIC 2+0 + space diversity, 2 × 56 MHz", 7.5, 2.4, 28, -65, "0.9 Gbps", 0.0046, 1.33)]
R001 = 22.0                                                              # mm/h exceeded 0.01 % of time (ITU-R P.837, north-west Jordan; desert is lower)
mw_rows, nlos = [], 0
for s in sites:
    if not s["backhaul"].startswith("Micro"): continue
    la, lo = float(s["lat"]), float(s["lon"])
    d2 = ((FL[:, 0] - la) * 110.6) ** 2 + ((FL[:, 1] - lo) * 111.3 * math.cos(math.radians(la))) ** 2
    j = int(np.argmin(d2)); d = math.sqrt(d2[j]); t = fib[j]
    mx, name, f, dish, tx, thr, cap, k, al = next(b for b in BANDS if d <= b[0])
    gain = 17.8 + 20 * math.log10(dish * f)
    fsl = 92.45 + 20 * math.log10(f) + 20 * math.log10(max(d, 0.1))
    rsl = tx + 2 * gain - fsl - 1.0
    gamma = k * R001 ** al; reff = 1 / (1 + d / (35 * math.exp(-0.015 * R001))); rain = gamma * d * reff
    margin = rsl - thr
    # terrain clearance: 60 % of first Fresnel zone, k = 4/3
    tt = np.linspace(0.02, 0.98, 60)
    zt = DEM.sample(la + (FL[j, 0] - la) * tt, lo + (FL[j, 1] - lo) * tt)
    za = float(DEM.sample(la, lo)) + float(s["antenna_height_agl_m"]) - 3; zb = float(DEM.sample(FL[j, 0], FL[j, 1])) + float(t["antenna_height_agl_m"]) - 3
    d1 = tt * d * 1000; d2m = (1 - tt) * d * 1000
    need = zt + d1 * d2m / (2 * 8.5e6) + 0.6 * np.sqrt(0.3 / f * d1 * d2m / (d * 1000))
    clr = float(np.min(za + (zb - za) * tt - need)); los = clr >= 0
    if not los: nlos += 1
    mw_rows.append([s["site_id"], t["site_id"], round(d, 1), name, cap, dish, round(gain, 1), round(fsl, 1), round(rsl, 1), round(margin, 1), round(rain, 1),
                    "OK" if margin >= rain else "rain-limited: larger dish / lower band", "LOS" if los else f"BLOCKED by {-clr:.0f} m: raise antennas or relay", round(clr)])
wcsv("microwave_links.csv", ["site", "far_end_fiber_site", "hop_km", "radio_configuration", "capacity", "dish_m", "antenna_gain_dBi", "free_space_loss_dB", "RSL_dBm", "fade_margin_dB",
     "rain_attenuation_0.01pct_dB", "availability_99.99", "terrain_line_of_sight", "clearance_m_vs_0.6F1"], mw_rows)
mw_cfg = Counter(r[3] for r in mw_rows); mw_rain = sum(1 for r in mw_rows if r[11] != "OK")

# =====================================================================================================
# Documents
# =====================================================================================================
wmd("LLD-01_RAN_Equipment_Design.md", f"""
# LLD-01 – RAN Equipment Design (vendor-neutral)
Low-level design of radio, baseband, antenna-line and site power equipment for all {fm(len(sites))} sites. Per-site result: `lld/ran_site_equipment.csv`; per-hub baseband pools: `lld/ran_hub_baseband.csv`. Unit capacities and power figures are typical 2026 carrier-grade values – replace with the chosen vendor's datasheet in `tools/lld_design.py` and re-run.

## 1. Equipment standards
{table(["Code", "Equipment", "Key specification", "Avg / peak power", "Network quantity"], [
 ["AAU-64", "n78 64T64R massive-MIMO AAU", "3400–3800 MHz, 400 MHz IBW / 300 MHz OBW, 320 W, 192 elements, 24 dBi, eCPRI 3 × 25GE, ≤ 30 kg, ≤ 0.5 m², IP65, −40…+55 °C", "900 / 1,300 W", fm(EQ["AAU64"])],
 ["R-8T", "n78 8T8R radio + 8-port antenna (rural)", "8 × 40 W, 300 MHz OBW, 17.5 dBi", "750 / 1,050 W", fm(EQ["R8T"])],
 ["RRU-L2", "n28 2T4R RRU (city)", "2 × 60 W, 703–748 / 758–803 MHz", "350 / 500 W", fm(EQ["RRU28_2T"])],
 ["RRU-L4", "n28 4T4R RRU (rural / highway)", "4 × 40 W", "550 / 750 W", fm(EQ["RRU28_4T"])],
 ["RRU-M", "n1 + n3 dual-band 4T4R RRU", "4 × 60 W shared, 1800 + 2100 MHz", "700 / 950 W", fm(EQ["RRU_MID"])],
 ["RRU-M3", "n3 4T4R RRU (suburban / highway)", "4 × 40 W", "450 / 620 W", fm(EQ["RRU_N3"])],
 ["ANT", "Passive antenna 2L4H + RET", "2 × 698–960 (15.5 dBi) + 4 × 1695–2690 MHz (17.5 dBi), 2.0 m city / 2.6 m rural, AISG 2.0 RET per band", "–", fm(EQ["ANT"])],
 ["DU", "Distributed unit (baseband)", "≥ 9 × 100 MHz 64T64R carriers + 12 FDD cells, 6 × 25GE fronthaul, 2 × 25GE midhaul, Class-C PTP, GNSS", "350 / 500 W", f"{fm(EQ['DU'])} at site + {fm(tot_du_pool)} pooled"],
 ["FHGW", "Fronthaul gateway / passive WDM mux (C-RAN site)", "12-λ MWDM/LWDM 25G, CPRI→eCPRI conversion for FDD RRUs, site alarms", "60 / 80 W", fm(EQ["FHGW"])],
 ["CSR", "Cell-site router (D-RAN site)", "4 × 25GE + 8 × 10GE, SR-MPLS/SRv6, Class-C boundary clock, SyncE", "120 / 150 W", fm(EQ["CSR"])]])}
Total NR cells: **{fm(cells_total)}** (city sector = 3 × n78 + n28 + n1 + n3; suburban without n1; rural = 3 × n78 + n28 + n1 + n3; highway = n28 + n3). Software licences follow cells, carriers, MIMO layers and CA combinations (n3 PCell + n1 + 3 × n78 DL CA; UL on n78 or FDD).

## 2. Standard site configurations
{table(["Configuration", "Sites", "Cells", "Avg load W", "Peak load W", "Rectifier kW (N+1)", "Li-ion battery kWh", "Grid kVA"], typ_rows)}
* DC power −48 V, 3 kW rectifier modules, N+1, sized for peak load + 0.2C battery recharge. Battery autonomy **4 h city, 8 h rural / highway** at average load (80 % depth of discharge). Li-ion with anti-theft lock and remote BMS.
* **{n_off} rural / highway sites without fibre or grid nearby are designed off-grid:** solar PV ≈ {max(x['pv'] for x in SE.values())} kWp (traffic-adaptive radio sleep, 5.5 peak-sun-hours) + battery + 15 kVA stand-by generator.
* Energy saving features mandatory in the tender: symbol / channel shutdown, carrier shutdown of 2nd–3rd n78 carrier at night, deep-sleep AAU – target ≥ 25 % energy reduction versus the averages above.
* Outdoor cabinets IP55 with heat exchanger (no air-conditioning) up to 50 °C ambient; Aqaba and Jordan Valley sites use the 55 °C variant with sun shield.

## 3. Baseband architecture
* **Dense urban / urban: C-RAN.** AAUs and RRUs connect over passive-WDM fronthaul (≤ 10 km, ≤ 75 µs one-way, 12-core drop) to the DU pool in their hub. Pool size = 0.75 DU per site (25 % pooling gain from non-coincident busy hours); inter-site CA, DL CoMP and coordinated scheduling run inside the pool – this is what lifts the cell-edge pixels to 1 Gbps.
* **Suburban / rural / highway: D-RAN.** DU in the site cabinet, 25GE midhaul on the access ring.
* **CU (CU-CP / CU-UP) virtualised** on the edge cloud at the six core / edge sites; one CU cluster per ≈ 500 sites; F1 over the aggregation network, latency budget ≤ 5 ms.
* **{len(hubs)} hubs**, of which {sum(1 for h in hubs if h['type'].startswith('C-RAN'))} are DU hotels with **{fm(tot_du_pool)} pooled DUs**; largest hub {max(r[5] for r in hub_rows)} DUs / {max(r[6] for r in hub_rows)} racks / {max(r[8] for r in hub_rows)} kW facility load. Hub rooms: 2N power feeds, N+1 cooling, 4 h battery + generator, GNSS antenna pair, fire suppression, access control.

## 4. Antenna line and installation standards
Jumpers ≤ 3 m, ½″ super-flex, 4.3-10 connectors, PIM ≤ −153 dBc; AISG RET daisy-chain per sector; AAU on independent bracket with ±15° mechanical azimuth and 0–10° tilt adjustment; wind load design 160 km/h; rooftop pole loading checked for 75 kg per sector (AAU + passive antenna + RRUs). GNSS receiver per D-RAN site and per hub; PTP as backup (and primary for C-RAN radios).

## 5. Parameter baseline (common to all vendors)
SCS 30 kHz (n78) / 15 kHz (FDD); TDD pattern DDDSU, special slot 10:2:2; SSB 8 beams n78; PCI from `output/sectors.csv`; digital tilt and RET per sector from the same file; max 4 DL layers per UE, MU-MIMO up to 16 layers; 256-QAM DL / 64-QAM UL (256-QAM UL where supported); n3 PCell with n78 SCells, UL Tx switching enabled; VoNR on n28/n3 with EVS-WB; inactivity timer 10 s; A3 offset 3 dB, TTT 320 ms (city) / 640 ms (rural).
""")

wmd("LLD-02_Core_Network_Design.md", f"""
# LLD-02 – 5G Core Network Design (vendor-neutral)
Cloud-native 5GC (3GPP Rel-17 SBA) sized for **{fm(SUBS)} subscribers** and **{fm(BH_GBPS)} Gbps busy-hour user plane**. Tables: `lld/core_nf_sizing.csv`, `lld/core_site_sizing.csv`.

## 1. Dimensioning inputs
{table(["Parameter", "Value"], [["Provisioned subscribers", fm(SUBS)], ["Simultaneously registered", f"{A['reg'] * 100:.0f} % → {fm(SUBS * A['reg'])}"], ["PDU sessions per registered sub", f"{A['sess']} (internet + IMS) → {fm(SUBS * A['reg'] * A['sess'])}"],
 ["Busy-hour user plane", f"{fm(BH_GBPS)} Gbps average, × {A['peak']} peak-to-mean"], ["VoNR subscribers", f"{A['volte'] * 100:.0f} % → {fm(erl)} Erlang (1 BHCA, {A['hold']} s)"],
 ["Design utilisation", f"{A['util'] * 100:.0f} % of rated capacity, N+1 per function, geo-redundant 1+1 between DC1 and DC2"]])}
Busy-hour signalling load (transactions per second): {', '.join(f'{k.replace("_", " ")} {fm(v)}' for k, v in tps.items())} → **{fm(sum(tps.values()))} procedures/s**, ≈ {fm(sum(tps.values()) * 6)} SBI transactions/s.

## 2. Network-function sizing (per data centre – each DC carries 100 % on failure of the other)
{table(["Network function", "Driver", "Unit", "Capacity / instance", "Instances (N+1)", "vCPU"], nf_rows)}
Control plane per DC: **{fm(vcpu_site)} vCPU** + 25 % platform overhead → **{cp_srv} servers** (2-socket, 96 usable vCPU, 512 GB RAM, 2 × 100GE). Subscriber database (UDR) synchronously replicated DC1↔DC2, asynchronous third replica in Irbid.

## 3. Site sizing (user plane distributed, control plane central)
{table(["Site", "UP share", "Peak Gbps", "UPF servers (N+1)", "CP servers", "Total servers", "Racks", "IT load kW", "400GE ports to transport"], core_rows)}
UPF server: 2 × 100GE SmartNIC, ≈ {A['upf_srv']} Gbps per server at 70 % load with full feature set (charging, LI, DPI-lite). UPF selection by TAC / DNAI so traffic breaks out at the nearest site; CDN caches and MEC hosts co-located on the N6 side of every UPF site.

## 4. Logical design
* **PLMN** 416-xx (MCC 416 Jordan; MNC from TRC). **TAC**: one per hub cluster ({len(hubs)} + rural TACs), TA lists sized ≤ 16 TACs to keep paging below {fm(tps['paging'])} /s.
* **Slices (S-NSSAI):** SST 1 / SD 000001 eMBB-1G (GFBR 1 Gbps tariff, 5QI 6/8/9 + dedicated GBR flow for the premium tier) · SST 1 / SD 000002 FWA · SST 1 / SD 000010 enterprise / private DNNs · SST 2 URLLC (future) · SST 1 / SD 0000FF public safety · O&M.
* **DNNs:** internet, ims, fwa, enterprise-<customer>, iot, sos.
* **Addressing:** IPv6-first. UE pools: one /40 per UPF site from the operator's RIPE NCC /29–/32 allocation (a /64 per UE session), IPv4 via NAT64/464XLAT and a CGNAT pool (≈ 1 public IPv4 per 64 users). Infrastructure: loopbacks 10.255.0.0/16, SBI 10.10.0.0/16 per DC, N2/N3 10.20.0.0/14 regionalised, O&M 10.250.0.0/16, all dual-stack with IPv6 ULA-free GUA on N3.
* **Naming / discovery:** 3GPP FQDNs (`amf1.set01.region01.amf.5gc.mnc0xx.mcc416.3gppnetwork.org`), NRF hierarchical (one per DC + PLMN-level), SCP model C for indirect communication, DNS anycast pairs in every site.
* **Voice:** IMS with VoNR only (no LTE layer in a pure-SA network) – EVS-WB, SRVCC not applicable; emergency calls on n28/n3; interconnect to other operators through SBC pair in DC1/DC2 (SIP-I / SIP), roaming through SEPP (N32) and IPX; devices without VoNR are a launch risk (see deployment risk register).
* **Security:** zones Untrust (N6/internet) – DMZ (SEPP, SBC, NEF) – Core SBA – O&M – LI; mutual TLS on all SBI with operator PKI; SUCI concealment (profile A) enabled; N2/N3 protected by IPsec from D-RAN sites / MACsec on owned fibre; firewalls + DDoS scrubbing on N6 (≥ 20 % of N6 capacity); LI (X1/X2/X3) mediation in both DCs.
* **Cloud platform:** Kubernetes CaaS on bare metal, SR-IOV / DPDK for UPF, separate clusters per zone, GitOps lifecycle, in-service upgrade with canary; observability (metrics, traces, logs) and NWDAF feeding the capacity trigger (sector PRB > 20 %).
* **Charging / policy:** converged CHF (online + offline), PCF rules for fair-use shaping that protects the single-user 1 Gbps experience; BSS integration over REST/TMF APIs.

## 5. Build phases
C0 facilities (M1–M2) → C1 control plane, UPF pool 1, IMS, first call (M2–M4) → C2 geo-redundancy tests, Amman Centre + Zarqa edge (M6–M9) → C3 Irbid and Aqaba regional DCs (M8–M12) → C4 UPF capacity steps of ≈ 400 Gbps per 600 accepted city sites (M13–M30) → C5 final dimensioning, DR drill, hand-over (M31–M36).
""")

lam_tot = sum(r[4] for r in dw_rows)
wmd("LLD-03_Transport_Network_Design.md", f"""
# LLD-03 – Transport Network Design (vendor-neutral)
Ports and router classes per node, optics, DWDM wavelength plan, IP/MPLS design and microwave link design. Tables: `transport_hub_ports.csv`, `transport_pop_ports.csv`, `transport_agg_ring_optics.csv`, `dwdm_wavelength_plan.csv`, `microwave_links.csv`.

## 1. Node classes
{table(["Class", "Where", "Ports", "Quantity"], [
 ["CSR", "every D-RAN site", "4 × 25GE + 8 × 10GE, Class-C BC", fm(EQ["CSR"])],
 ["FHGW / passive WDM", "every C-RAN site", "12 λ × 25G on one fibre pair (+ protection pair)", fm(EQ["FHGW"])]] +
 [[k.split(" ")[0], "hub (pair per hub)", k[k.index("(") + 1:-1], f"{v} hubs → {2 * v} routers"] for k, v in hub_cls.items()] +
 [[k.split(" ")[0], "PoP / core site (pair)", k[k.index("(") + 1:-1], f"{v} PoPs → {2 * v} routers"] for k, v in Counter(r[4] for r in pop_rows).items()] +
 [["DWDM", "backbone nodes", "ROADM / FOADM, 400G coherent line cards", f"{len(dw_rows)} routes, {lam_tot} wavelengths day 1 (400G on rings, 100G on light spurs)"]])}

## 2. Access and aggregation dimensioning
* **Access ring (D-RAN), 25GE:** ≤ 8 sites, dual-homed on the hub router pair. Design load = 1.5 × busy-hour mean + one sector at its 5.2 Gbps peak. Result over {len(dran_rings)} D-RAN rings: mean {np.mean(dran_rings):.1f} Gbps, max {max(dran_rings):.1f} Gbps → **≤ {100 * max(dran_rings) / 25:.0f} % of 25GE**; upgrade path 50GE optics on the same fibre.
* **C-RAN access ring:** physical fibre ring carrying point-to-point fronthaul wavelengths (9 × 25G eCPRI for the AAUs + 3 × 10G for FDD RRUs per site, passive WDM, both directions of the ring for protection). Optical budget: 25G LR ≤ 10 km, ≤ 6.3 dB path + 2 × 2.5 dB mux.
* **Hub routers:** one end of every access ring per router + half of the pooled DUs + 100GE uplinks; peak hub load max {max(r[9] for r in port_rows)} Gbps, uplink utilisation at peak ≤ {max(int(r[10][:-2]) for r in port_rows)} %.
* **Aggregation rings, 100GE:** ≤ 5 hubs per ring; optic by mean span: {', '.join(f'{v} × {k}' for k, v in Counter(r[5] for r in agg_rows).items())}. Ring load = sum of hub peaks × 0.7 diversity; rings above 60 Gbps are split or moved to 2 × 100GE LAG (check column peak_Gbps per PoP).
* **PoP routers:** aggregation rings on 100GE, 400GE to DWDM / core; busiest PoP {pop_rows[0][0]} {pop_rows[0][5]} Gbps busy hour.

## 3. DWDM backbone – wavelength plan (day 1, 400G coherent, 75 GHz flex-grid, 64 slots per fibre pair)
{table(["Route", "Ring", "km", "Design demand Gbps", "λ day 1", "Channel plan", "In-line amplifier sites", "Node type", "Line optics"], dw_rows)}
Demand = busy-hour traffic of the clusters homed on the route × 1.3 peak × 1.5 growth; ring routes carry the **whole ring demand** so any single fibre cut is survivable (OTN 1+1 / ROADM restoration < 50 ms for protected services). The southern ring also carries ≈ {fm(INTL)} Gbps of international transit from the Aqaba cable landing (40 % of internet volume, after 60 % in-country cache hit). EDFA every ≤ 80 km (Raman on > 100 km desert spans), G.652.D fibre, span loss budget 0.22 dB/km + 0.5 dB per splice-km allowance; OSNR ≥ 24 dB end-to-end for 400G 16-QAM. Fill after 3 years ≈ {100 * lam_tot / (len(dw_rows) * 64):.0f} % of slots – ample headroom.

## 4. IP / MPLS design
* **IGP:** IS-IS L2, one instance per domain (metro Amman, north, south, east), multi-instance at PoP border routers; **SR-MPLS with TI-LFA (< 50 ms)**, SRv6-ready hardware; BFD 3 × 10 ms on ring links.
* **Services:** BGP EVPN / L3VPN per function – RAN-CP (N2), RAN-UP (N3), F1 (C-RAN midhaul), O&M, Sync, Enterprise; route reflectors in DC1 / DC2 / Irbid; inter-domain with BGP-LU + seamless MPLS.
* **Slicing:** FlexE / flex-algo hard slices for eMBB-1G (low-delay flex-algo), FWA, enterprise, O&M.
* **QoS (8 classes):** PTP/sync EF-strict · network control CS6 · RAN control + VoNR signalling CS5 · VoNR media EF · eMBB-1G GFBR AF41 · FWA / premium data AF31 · best-effort data BE · O&M AF21; 5QI→DSCP mapping done in DU/CU and UPF; ring links shaped at 95 %.
* **MTU 9,100** end-to-end (GTP-U + SRv6 headroom); IPv6 on all infrastructure links, IPv4 only where a legacy management system needs it.
* **Addressing:** loopbacks 10.255.<domain>.0/24, p2p /31 from 10.254.0.0/16, IPv6 /127 from the infrastructure /40; site subnets summarised per hub.
* **Synchronisation:** ePRTC + GNSS at DC1, DC2, Irbid, Aqaba; PRTC-B at every hub; G.8275.1 full on-path support, every router Class-C boundary clock; budget ±1.5 µs network, ±130 ns relative inside a C-RAN cluster; SyncE everywhere; GNSS holdover 24 h at hubs.
* **Security:** MACsec on owned fibre rings, IPsec (N2/N3) from sites on leased or microwave links, control-plane policing, RTBH / flowspec, out-of-band management network via LTE-free DCN on dedicated VPN.

## 5. Microwave design ({len(mw_rows)} rural / highway sites)
{table(["Radio configuration", "Hops"], sorted(mw_cfg.items(), key=lambda q: -q[1]))}
* Per-hop budget in `microwave_links.csv`: free-space loss, antenna gain, receive level, fade margin, ITU-R P.530/P.838 rain attenuation at R0.01 = {R001:.0f} mm/h, and a **terrain line-of-sight check on the Copernicus DEM** (60 % of the first Fresnel zone, k = 4/3, antennas 3 m below tower top).
* Result: **{len(mw_rows) - nlos} hops clear, {nlos} blocked by terrain** – those need a taller structure, a relay / passive repeater, or a different far end (column `terrain_line_of_sight` gives the missing clearance); {mw_rain} hops are rain-limited and take the next dish size.
* Hops longer than 16 km deliver ≤ 1.3 Gbps: those sites start with one 100 MHz n78 carrier on the rural 8T8R radio and move to fibre or a second hop when traffic requires.
* Frequency coordination and licensing with TRC per hop; E-band under light licensing where available; XPIC with ACM down to QPSK for 99.999 % of the control / voice traffic class.
""")

# workbook
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); first = True
for fcsv in ["ran_site_equipment", "ran_hub_baseband", "core_nf_sizing", "core_site_sizing", "transport_hub_ports", "transport_pop_ports", "transport_agg_ring_optics", "dwdm_wavelength_plan", "microwave_links"]:
    ws = wb.active if first else wb.create_sheet(); first = False; ws.title = fcsv[:31]
    for i, r in enumerate(csv.reader(open(os.path.join(LLD, fcsv + ".csv"), encoding="utf-8"))):
        ws.append([float(c) if i and c.replace(".", "", 1).replace("-", "", 1).isdigit() else c for c in r])
    for c in ws[1]: c.fill, c.font, c.alignment = PatternFill("solid", fgColor="1F3A5F"), Font(bold=True, color="FFFFFF"), Alignment(wrap_text=True)
    ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = 20
wb.save(os.path.join(LLD, "Jordan_5G_LLD.xlsx"))
print("RAN", dict(EQ), "cells", cells_total, "pooled DU", tot_du_pool, "offgrid", n_off)
print("CORE vcpu/DC", vcpu_site, "cp servers", cp_srv, [(r[0], r[5], r[6]) for r in core_rows])
print("TRANSPORT hub classes", dict(hub_cls), "dran ring max", round(max(dran_rings), 1), "DWDM λ", lam_tot, "MW", len(mw_rows), "blocked", nlos, "rain", mw_rain, dict(mw_cfg))
